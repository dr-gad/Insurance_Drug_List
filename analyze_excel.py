import sys
import io
import json
import re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook('لستة ادويه التأمين.xlsx')

# ==== Sheet 1: المجاني ====
ws = wb['المجاني']
all_rows = []
for i, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
    non_none = [v for v in row if v is not None]
    if non_none:
        data = [str(v) if v is not None else None for v in row[:10]]
        all_rows.append({'row': i, 'data': data})

print("=== Total rows in المجاني:", len(all_rows))
print()

# Print groups and subgroups
for r in all_rows:
    data = r['data']
    non_null = [d for d in data if d]
    first = non_null[0] if non_null else ''
    if re.match(r'Group\s*\(', first, re.IGNORECASE):
        print(f"GROUP: row {r['row']} => {first}")
    elif len(non_null) == 1 and not str(first).strip().replace('.','').isdigit():
        if len(first) > 5:  # avoid short junk
            print(f"  SUBGROUP: row {r['row']} => {first}")

print()
print("=== Last 30 rows:")
for r in all_rows[-30:]:
    print(f"  row {r['row']}: {r['data']}")

print()

# ==== Sheet 2: تجاري ====
ws2 = wb['تجاري']
print("=== Sheet تجاري ===")
for i, row in enumerate(ws2.iter_rows(min_row=1, values_only=True), start=1):
    non_none = [v for v in row if v is not None]
    if non_none:
        print(f"  row {i}: {[str(v) for v in row if v is not None]}")

print()
# ==== Sheet 3: ورقة1 ====
ws3 = wb['ورقة1']
print("=== Sheet ورقة1 ===")
for i, row in enumerate(ws3.iter_rows(min_row=1, values_only=True), start=1):
    non_none = [v for v in row if v is not None]
    if non_none:
        print(f"  row {i}: {[str(v) for v in row if v is not None]}")
