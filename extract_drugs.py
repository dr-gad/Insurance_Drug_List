import sys
import io
import json
import re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook('لستة ادويه التأمين.xlsx')

def clean(val):
    if val is None:
        return None
    s = str(val).strip()
    # Remove non-breaking spaces
    s = s.replace('\xa0', ' ')
    # normalize whitespace
    s = re.sub(r'\s+', ' ', s)
    return s if s else None

def normalize_authority(raw):
    """Normalize authority value preserving all info."""
    if not raw:
        return None, None

    r = raw.strip()
    r_lower = r.lower()

    # Detect committee tags
    committee = None
    if 'لجنة ms' in r_lower or 'لجنه ms' in r_lower:
        committee = 'لجنة MS'
    elif 'لجنة القلب' in r_lower or 'لجنه القلب' in r_lower:
        committee = 'لجنة القلب'
    elif 'لجنة عليا' in r_lower or 'لجنه عليا' in r_lower:
        committee = 'لجنة عليا'

    # Normalize base authority
    if r_lower.startswith('gp'):
        base = 'gp'
    elif r_lower.startswith('con') or 'con' in r_lower:
        base = 'con'
    elif r_lower.startswith('cons'):
        base = 'con'
    elif r_lower.startswith('sp') or r_lower == 'sp':
        base = 'sp'
    elif committee:
        # Arabic-only authority
        base = 'con'  # treated as consultant-level
    else:
        base = r_lower

    return base, committee

# =============================================
# Parse the main sheet: المجاني
# =============================================
ws = wb['المجاني']
free_drugs = []

current_group_num  = None
current_group_name = None
current_sub_a      = None   # e.g. "A- Antibiotics", "B-Anti -Tuberculous Drugs"
current_subgroup   = None   # e.g. "Benzyl penicillin...", "Ampicillin Derivatives"

# Rows that are pure section headers (no serial, one cell filled)
# We need to distinguish:
#  - Group (N) rows
#  - Group name rows (directly after Group (N))
#  - Sub-A rows  (e.g. "A- Antibiotics", "B-Anti-Tuberculous", "C-...", "D-...")
#  - Sub-subgroup rows (e.g. "Benzyl penicillin...", "Ampicillin Derivatives")
#  - Arabic section headers: "ادوية اللجنة العليا الاورام", "ادوية اللجنة العليا للدواء"
# Drug rows: first col is integer serial

def is_group_row(first):
    return bool(re.match(r'Group\s*\(?\d+\)?', first, re.IGNORECASE))

def is_sub_section(first):
    """Top-level lettered sub-section inside a group e.g. A- Antibiotics"""
    return bool(re.match(r'^[A-Z]\s*[-–]', first))

def is_arabic_section(first):
    """Pure Arabic section headers like 'ادوية اللجنة العليا الاورام'"""
    # Contains Arabic chars and no Latin digits at start
    has_arabic = bool(re.search(r'[\u0600-\u06FF]', first))
    starts_digit = bool(re.match(r'^\d', first))
    return has_arabic and not starts_digit

for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    cols = [clean(c) for c in row]
    non_null = [c for c in cols if c]
    if not non_null:
        continue

    first = non_null[0]

    # ── Group (N) header ──────────────────────────────────────
    m = re.match(r'Group\s*\(?(\d+)\)?', first, re.IGNORECASE)
    if m:
        current_group_num  = int(m.group(1))
        current_group_name = None
        current_sub_a      = None
        current_subgroup   = None
        continue

    # ── Arabic section headers (groups 34/35) ────────────────
    if is_arabic_section(first) and len(non_null) == 1:
        # Check if it's one of the known Arabic group headers
        if 'اللجنة العليا الاورام' in first or 'اللجنه العليا الاورام' in first:
            current_group_num  = 34
            current_group_name = 'أدوية اللجنة العليا للأورام'
            current_sub_a      = None
            current_subgroup   = None
        elif 'اللجنة العليا للدواء' in first or 'اللجنه العليا للدواء' in first:
            current_group_num  = 35
            current_group_name = 'أدوية اللجنة العليا للدواء'
            current_sub_a      = None
            current_subgroup   = None
        else:
            # Just a subgroup label in Arabic
            current_subgroup = first
        continue

    # ── Single-cell text rows ─────────────────────────────────
    if len(non_null) == 1 and not re.match(r'^\d+$', first):
        if current_group_name is None and current_group_num is not None:
            # This is the group name (first text row after Group (N))
            current_group_name = first
            current_sub_a      = None
            current_subgroup   = None
        elif is_sub_section(first):
            # e.g. "A- Antibiotics", "B-Anti-Tuberculous"
            current_sub_a  = first
            current_subgroup = None
        else:
            # subgroup within a section
            current_subgroup = first
        continue

    # ── Drug row ─────────────────────────────────────────────
    serial = cols[0]
    if serial and re.match(r'^\d+$', str(serial).strip()):
        drug_name = cols[1] if len(cols) > 1 else None
        conc      = cols[2] if len(cols) > 2 else None
        drug_form = cols[3] if len(cols) > 3 else None
        auth_raw  = cols[4] if len(cols) > 4 else None

        auth_base, committee = normalize_authority(auth_raw)

        # Build subgroup string
        sub_parts = []
        if current_sub_a:    sub_parts.append(current_sub_a)
        if current_subgroup: sub_parts.append(current_subgroup)
        subgroup_str = ' > '.join(sub_parts) if sub_parts else None

        entry = {
            'serial':        int(serial),
            'drug_name':     drug_name,
            'concentration': conc,
            'drug_form':     drug_form,
            'authority_raw': auth_raw,
            'authority':     auth_base,
            'committee':     committee,
            'group_num':     current_group_num,
            'group_name':    current_group_name,
            'subgroup':      subgroup_str,
            'sheet':         'free',
        }
        free_drugs.append(entry)

print(f"Free drugs: {len(free_drugs)}")

# Check Group 3 samples
print("\nGroup 3 samples:")
for d in free_drugs:
    if d['group_num'] == 3:
        print(f"  #{d['serial']} | {d['drug_name']} | sub={d['subgroup']}")
        if d['serial'] > 107:
            break

# Check groups 34, 35
print("\nGroup 34 samples:")
for d in free_drugs:
    if d['group_num'] == 34:
        print(f"  #{d['serial']} | {d['drug_name']}")
        
print("\nGroup 35 samples:")
for d in free_drugs:
    if d['group_num'] == 35:
        print(f"  #{d['serial']} | {d['drug_name']}")

# Check authority
print("\nCommittee samples:")
for d in free_drugs:
    if d['committee']:
        print(f"  #{d['serial']} | {d['drug_name']} | auth={d['authority_raw']} | committee={d['committee']}")

# Groups overview
from collections import Counter
grp_counts = Counter(d['group_num'] for d in free_drugs)
print(f"\nGroup counts:")
groups_map = {}
for d in free_drugs:
    if d['group_num'] not in groups_map:
        groups_map[d['group_num']] = d['group_name']
for k in sorted(groups_map.keys()):
    print(f"  Group {k} ({grp_counts[k]}): {groups_map[k]}")

# =============================================
# Sheet 2: تجاري
# =============================================
ws2 = wb['تجاري']
commercial_drugs = []

for i, row in enumerate(ws2.iter_rows(min_row=2, values_only=True), start=2):
    cols = [clean(c) for c in row]
    unit1 = cols[0] if len(cols) > 0 else None
    name1 = cols[1] if len(cols) > 1 else None
    unit2 = cols[2] if len(cols) > 2 else None
    name2 = cols[3] if len(cols) > 3 else None
    if name1:
        commercial_drugs.append({'drug_name': name1, 'unit': unit1, 'sheet': 'commercial'})
    if name2:
        commercial_drugs.append({'drug_name': name2, 'unit': unit2, 'sheet': 'commercial'})

print(f"\nCommercial drugs: {len(commercial_drugs)}")

# =============================================
# Sheet 3: ورقة1
# =============================================
ws3 = wb['ورقة1']
special_drugs = []

for i, row in enumerate(ws3.iter_rows(min_row=1, values_only=True), start=1):
    cols = [clean(c) for c in row]
    non_null = [c for c in cols if c]
    if non_null:
        raw = non_null[0]
        m = re.match(r'^(.*?)\s*-\s*(.*?)\s*-\s*\[(.*?)\]$', raw)
        if m:
            special_drugs.append({
                'drug_name':  m.group(1).strip(),
                'drug_form':  m.group(2).strip(),
                'trade_name': m.group(3).strip(),
                'sheet':      'special',
            })
        else:
            special_drugs.append({'drug_name': raw, 'sheet': 'special'})

print(f"Special drugs: {len(special_drugs)}")

# =============================================
# Save
# =============================================
data = {
    'free_drugs':       free_drugs,
    'commercial_drugs': commercial_drugs,
    'special_drugs':    special_drugs,
}

with open('drugs_data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print("\nSaved drugs_data.json")
total = len(free_drugs) + len(commercial_drugs) + len(special_drugs)
print(f"Total: {total}")
